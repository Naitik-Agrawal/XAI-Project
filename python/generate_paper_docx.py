import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def clean_math_text(math_str):
    """Converts LaTeX math snippets into clean Unicode mathematical symbols."""
    s = math_str
    replacements = [
        (r"\\mathbf\{CM\}", "CM"),
        (r"\\mathbf\{x\}", "x"),
        (r"\\mathbf\{X\}", "X"),
        (r"\\mathbb\{R\}\^\{14\}", "ℝ¹⁴"),
        (r"\\mathbb\{R\}", "ℝ"),
        (r"\\mathbb\{E\}", "E"),
        (r"\\mathbb\{Z\}", "ℤ"),
        (r"\\mathcal\{O\}", "𝒪"),
        (r"\\mathcal\{P\}", "𝒫"),
        (r"\\mathcal\{T\}_1", "𝒯₁"),
        (r"\\mathcal\{T\}_2", "𝒯₂"),
        (r"\\mathcal\{T\}_3", "𝒯₃"),
        (r"\\mathcal\{T\}_4", "𝒯₄"),
        (r"\\mathcal\{T\}", "𝒯"),
        (r"\\mathcal\{A\}", "𝒜"),
        (r"\\phi_j", "ϕⱼ"),
        (r"\\phi", "ϕ"),
        (r"\\Delta t", "Δt"),
        (r"\\Delta", "Δ"),
        (r"\\le", "≤"),
        (r"\\ge", "≥"),
        (r"\\to", "→"),
        (r"\\sum", "∑"),
        (r"\\times", "×"),
        (r"\\cdot", "·"),
        (r"\\dots", "..."),
        (r"\\in", "∈"),
        (r"\\subset", "⊂"),
        (r"\\infty", "∞"),
        (r"\\mid", "|"),
        (r"\\text\{ms\}", " ms"),
        (r"\\text\{min\}", "min"),
        (r"\\text\{max\}", "max"),
        (r"\\text\{TN\}", "TN"),
        (r"\\text\{TP\}", "TP"),
        (r"\\text\{FN\}", "FN"),
        (r"\\text\{FP\}", "FP"),
        (r"\\text\{FPR\}", "FPR"),
        (r"\\text\{FNR\}", "FNR"),
        (r"\\text\{ALLOW\}", "ALLOW"),
        (r"\\text\{RESTRICT\}", "RESTRICT"),
        (r"\\text\{VERIFY_QUARANTINE\}", "VERIFY_QUARANTINE"),
        (r"\\text\{VERIFY\_QUARANTINE\}", "VERIFY_QUARANTINE"),
        (r"\\text\{BLOCK\}", "BLOCK"),
        (r"\\text\{", ""),
        (r"\}", ""),
        (r"\{", ""),
        (r"\\", ""),
    ]
    for pattern, repl in replacements:
        s = re.sub(pattern, repl, s)
    return s

def add_formatted_runs(paragraph, text, base_font="Times New Roman", base_size=Pt(11), base_color=RGBColor(35, 35, 35)):
    """Parses markdown inline bold, italic, code, and math tokens into runs."""
    token_pattern = re.compile(
        r'(`(?P<code>[^`]+)`)'
        r'|(\$(?P<math>[^\$]+)\$)'
        r'|(\*\*(?P<bold>[^\*]+)\*\*)'
        r'|(\*(?P<italic>[^\*]+)\*)'
    )
    
    pos = 0
    for match in token_pattern.finditer(text):
        start, end = match.span()
        if start > pos:
            plain = text[pos:start]
            r = paragraph.add_run(plain)
            r.font.name = base_font
            r.font.size = base_size
            r.font.color.rgb = base_color
            
        group_dict = match.groupdict()
        if group_dict.get("code"):
            r = paragraph.add_run(group_dict["code"])
            r.font.name = "Consolas"
            r.font.size = Pt(base_size.pt * 0.92)
            r.font.color.rgb = RGBColor(160, 40, 40)
        elif group_dict.get("math"):
            math_clean = clean_math_text(group_dict["math"])
            r = paragraph.add_run(math_clean)
            r.font.name = "Cambria Math"
            r.font.size = base_size
            r.font.italic = True
            r.font.color.rgb = RGBColor(15, 35, 80)
        elif group_dict.get("bold"):
            r = paragraph.add_run(group_dict["bold"])
            r.font.name = base_font
            r.font.size = base_size
            r.font.bold = True
            r.font.color.rgb = base_color
        elif group_dict.get("italic"):
            r = paragraph.add_run(group_dict["italic"])
            r.font.name = base_font
            r.font.size = base_size
            r.font.italic = True
            r.font.color.rgb = base_color
            
        pos = end
        
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.font.name = base_font
        r.font.size = base_size
        r.font.color.rgb = base_color

def set_cell_background(cell, hex_color):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, top_color="1F3A60", bottom_color="1F3A60", inside_h_color="E2E8F0"):
    """Applies clean academic borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="12" w:space="0" w:color="{top_color}"/>
        <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{bottom_color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{inside_h_color}"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>''')
    tblPr.append(borders)

def set_cell_margins(table, top=100, bottom=100, left=140, right=140):
    """Sets cell padding for all cells in table (in dxa)."""
    tblPr = table._tbl.tblPr
    tblCellMar = parse_xml(f'''<w:tblCellMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tblCellMar>''')
    tblPr.append(tblCellMar)

def format_academic_table(doc, header, data_rows, col_widths, font_size=Pt(9)):
    """Builds and formats a professional IEEE/ACM-style table."""
    num_rows = len(data_rows) + 1
    num_cols = len(header)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_margins(table, top=80, bottom=80, left=110, right=110)

    # Header Row
    hdr_cells = table.rows[0].cells
    for col_idx, heading_text in enumerate(header):
        cell = hdr_cells[col_idx]
        set_cell_background(cell, "1F3A60")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        clean_hdr = heading_text.replace("**", "").replace("`", "")
        r = p.add_run(clean_hdr)
        r.font.name = "Times New Roman"
        r.font.size = font_size
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8F9FA" if (row_idx % 2 == 1) else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            if col_idx >= len(row_cells):
                continue
            cell = row_cells[col_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            val_clean = cell_value.strip()
            # Centering logic for numerical / status cells
            if any(token in val_clean for token in ["%", "ms", "s", "MB/s", "KB", "0.", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9."]) and len(val_clean) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif col_idx == 0 and len(val_clean) < 5:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            add_formatted_runs(p, cell_value, base_font="Times New Roman", base_size=font_size, base_color=RGBColor(30, 30, 30))

    # Apply column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)

    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def add_display_equation(doc, math_latex, eq_number=""):
    """Adds a beautifully formatted mathematical display block."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    clean_eq = clean_math_text(math_latex)
    r = p.add_run(clean_eq)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = RGBColor(10, 30, 70)
    
    if eq_number:
        r_num = p.add_run(f"    ({eq_number})")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(10)
        r_num.font.italic = False
        r_num.font.color.rgb = RGBColor(110, 110, 110)

def add_image_figure(doc, img_path, caption_text, width=Inches(5.9)):
    """Inserts an image centered on page with italic caption below."""
    if not os.path.exists(img_path):
        print(f"WARNING: Image not found: {img_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)
    
    # Caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.line_spacing = 1.1
    
    add_formatted_runs(p_cap, caption_text, base_font="Times New Roman", base_size=Pt(9.5), base_color=RGBColor(60, 60, 60))

def parse_markdown_paper(md_path):
    """Parses markdown file into structured blocks."""
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()

    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        sline = line.strip()

        if not sline:
            i += 1
            continue

        if sline == "---":
            blocks.append({"type": "divider"})
            i += 1
            continue

        if sline.startswith("# "):
            blocks.append({"type": "title", "content": sline[2:]})
            i += 1
            continue

        if sline.startswith("## "):
            blocks.append({"type": "h1", "content": sline[3:]})
            i += 1
            continue

        if sline.startswith("### "):
            blocks.append({"type": "h2", "content": sline[4:]})
            i += 1
            continue

        if sline.startswith("![](") or sline.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", sline)
            if m:
                alt = m.group(1)
                src = m.group(2)
                caption = ""
                if i + 1 < len(lines) and lines[i+1].strip().startswith("*Figure"):
                    caption = lines[i+1].strip()
                    i += 1
                blocks.append({"type": "image", "alt": alt, "src": src, "caption": caption})
                i += 1
                continue

        if sline.startswith("```"):
            lang = sline[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "lang": lang, "content": "\n".join(code_lines)})
            continue

        if sline.startswith("$$"):
            math_lines = []
            if len(sline) > 2 and not sline.endswith("$$"):
                math_lines.append(sline[2:])
            elif sline.endswith("$$") and len(sline) > 4:
                math_lines.append(sline[2:-2])
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    math_lines.append(lines[i])
                    i += 1
                if i < len(lines) and lines[i].strip().endswith("$$"):
                    end_l = lines[i].strip()[:-2]
                    if end_l:
                        math_lines.append(end_l)
            blocks.append({"type": "display_math", "content": "\n".join(math_lines).strip()})
            i += 1
            continue

        if sline.startswith("|"):
            tbl_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            blocks.append({"type": "table", "content": tbl_lines})
            continue

        if sline.startswith("- ") or sline.startswith("* "):
            blocks.append({"type": "bullet", "content": sline[2:]})
            i += 1
            continue

        m_num = re.match(r"^(\d+)\.\s+(.*)", sline)
        if m_num:
            blocks.append({"type": "numbered", "num": m_num.group(1), "content": m_num.group(2)})
            i += 1
            continue

        # Regular paragraph
        para_lines = [line]
        i += 1
        while i < len(lines):
            nl = lines[i]
            snl = nl.strip()
            if not snl:
                break
            if (snl.startswith("#") or snl.startswith("---") or snl.startswith("![") or 
                snl.startswith("```") or snl.startswith("$$") or snl.startswith("|") or 
                snl.startswith("- ") or snl.startswith("* ") or re.match(r"^(\d+)\.\s+", snl)):
                break
            para_lines.append(nl)
            i += 1
        blocks.append({"type": "paragraph", "content": " ".join([l.strip() for l in para_lines])})

    return blocks

def generate_paper_docx(md_path, output_docx_path):
    """Main compilation routine generating the complete academic Word manuscript."""
    print("Parsing research paper markdown...")
    blocks = parse_markdown_paper(md_path)
    print(f"Loaded {len(blocks)} structured blocks.")

    doc = docx.Document()

    # 1. Page Setup: A4, 1-inch margins
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    # 2. Header: Running title
    header = sec.header
    h_p = header.paragraphs[0]
    h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = h_p.add_run("XAI-ZT-P2P: Explainable AI-Driven Zero-Trust Framework for P2P Systems")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(130, 130, 130)

    # 3. Footer: Page X of NUMPAGES
    footer = sec.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_p.paragraph_format.line_spacing = 1.0
    r_f1 = f_p.add_run("Page ")
    r_f1.font.name = "Times New Roman"
    r_f1.font.size = Pt(9)
    r_f1.font.color.rgb = RGBColor(120, 120, 120)

    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr1 = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r_f1._r.append(fld1)
    r_f1._r.append(instr1)
    r_f1._r.append(fld2)
    r_f1._r.append(fld3)

    r_f2 = f_p.add_run(" of ")
    r_f2.font.name = "Times New Roman"
    r_f2.font.size = Pt(9)
    r_f2.font.color.rgb = RGBColor(120, 120, 120)

    fld4 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr2 = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fld5 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld6 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r_f2._r.append(fld4)
    r_f2._r.append(instr2)
    r_f2._r.append(fld5)
    r_f2._r.append(fld6)

    # State tracking
    current_section = ""
    table_index = 0
    math_eq_index = 0

    # Custom column widths for the 7 tables
    table_col_widths = {
        0: [1.25, 2.50, 2.52],                                      # Table I: 3 cols
        1: [0.30, 1.35, 1.25, 3.37],                                # Table II: 4 cols
        2: [0.95, 0.75, 0.65, 0.70, 1.25, 1.97],                    # Table III: 6 cols
        3: [0.95, 0.50, 0.50, 0.48, 0.48, 0.50, 0.45, 0.52, 0.50, 0.52, 0.87], # Table IV: 11 cols
        4: [1.85, 1.05, 1.15, 0.95, 1.27],                          # Table V: 5 cols
        5: [0.25, 0.95, 1.15, 0.45, 0.80, 0.50, 0.55, 0.95, 0.67], # Table VI: 9 cols
        6: [1.35, 2.45, 2.47]                                       # Table VII: 3 cols
    }

    # Equation numbers mapping
    eq_numbers = ["1", "2", "3", "4", "5", "6"]

    # Iterate through blocks
    for b_idx, block in enumerate(blocks):
        btype = block["type"]

        if btype == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(block["content"])
            r.font.name = "Times New Roman"
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 54, 93)

        elif btype == "paragraph" and ("Authors:" in block["content"] or "Target Publication Venue:" in block["content"]):
            # Document metadata header
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Author line
            r_author = p.add_run("Naitik Agrawal\n")
            r_author.font.name = "Times New Roman"
            r_author.font.size = Pt(12)
            r_author.font.bold = True
            r_author.font.color.rgb = RGBColor(35, 35, 35)

            r_affil = p.add_run("Advanced Cybersecurity & Networking Research Laboratory\nResearch Laboratory for Decentralized Systems and Network Security\nEmail: itsnaitik18@gmail.com\n")
            r_affil.font.name = "Times New Roman"
            r_affil.font.size = Pt(10)
            r_affil.font.italic = True
            r_affil.font.color.rgb = RGBColor(70, 70, 70)

            # Metadata callout
            meta_tbl = doc.add_table(rows=1, cols=1)
            meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            m_cell = meta_tbl.cell(0, 0)
            m_cell.width = Inches(6.27)
            set_cell_background(m_cell, "F0F4F8")
            
            tcPr = m_cell._tc.get_or_add_tcPr()
            b = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>
            </w:tcBorders>''')
            tcPr.append(b)

            mar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
                <w:top w:w="100" w:type="dxa"/>
                <w:bottom w:w="100" w:type="dxa"/>
                <w:left w:w="160" w:type="dxa"/>
                <w:right w:w="160" w:type="dxa"/>
            </w:tcMar>''')
            tcPr.append(mar)

            mp = m_cell.paragraphs[0]
            mp.paragraph_format.space_before = Pt(2)
            mp.paragraph_format.space_after = Pt(2)
            mp.paragraph_format.line_spacing = 1.15
            
            meta_text = (
                "**Target Publication Venue:** IEEE Transactions on Network and Service Management / ACM CCS\n"
                "**Artifact Repository & Empirical Prototype:** `https://github.com/research-lab/xai-zt-p2p`\n"
                "**Classification Schema:** Network Security; Peer-to-Peer Systems; Explainable AI; Adaptive Transport Protocols"
            )
            for mline in meta_text.splitlines():
                mlp = m_cell.add_paragraph() if mp.text else mp
                mlp.paragraph_format.space_before = Pt(1)
                mlp.paragraph_format.space_after = Pt(1)
                mlp.paragraph_format.line_spacing = 1.1
                add_formatted_runs(mlp, mline, base_font="Times New Roman", base_size=Pt(9.5), base_color=RGBColor(45, 45, 45))

            # Spacer
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        elif btype == "h1":
            current_section = block["content"]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            
            r = p.add_run(block["content"])
            r.font.name = "Times New Roman"
            r.font.size = Pt(13.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 54, 93)

        elif btype == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            r = p.add_run(block["content"])
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(43, 108, 176)

        elif btype == "paragraph":
            content = block["content"]
            
            # Check if this paragraph is a table title
            if content.startswith("**TABLE ") and "**" in content[8:]:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(10), base_color=RGBColor(27, 54, 93))
                continue

            # Check if Abstract / Keywords paragraph
            if current_section == "Abstract":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.25)
                add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(10), base_color=RGBColor(40, 40, 40))
                continue

            # Standard body paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(11), base_color=RGBColor(35, 35, 35))

        elif btype == "bullet":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.35)
            
            # Bullet symbol
            r_bullet = p.add_run("•  ")
            r_bullet.font.name = "Times New Roman"
            r_bullet.font.size = Pt(11)
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(43, 108, 176)
            
            add_formatted_runs(p, block["content"], base_font="Times New Roman", base_size=Pt(11), base_color=RGBColor(35, 35, 35))

        elif btype == "numbered":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            num_str = block["num"]
            content = block["content"]
            
            # If in References section, format as academic citation
            if current_section == "References":
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.first_line_indent = Inches(-0.4)
                r_num = p.add_run(f"[{num_str}]  ")
                r_num.font.name = "Times New Roman"
                r_num.font.size = Pt(10)
                r_num.font.bold = True
                r_num.font.color.rgb = RGBColor(27, 54, 93)
                add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(10), base_color=RGBColor(40, 40, 40))
            else:
                p.paragraph_format.left_indent = Inches(0.35)
                r_num = p.add_run(f"{num_str}.  ")
                r_num.font.name = "Times New Roman"
                r_num.font.size = Pt(11)
                r_num.font.bold = True
                r_num.font.color.rgb = RGBColor(43, 108, 176)
                add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(11), base_color=RGBColor(35, 35, 35))

        elif btype == "table":
            tbl_lines = block["content"]
            rows = [l.strip() for l in tbl_lines if l.strip()]
            if len(rows) >= 3:
                header = [c.strip() for c in rows[0].split('|')[1:-1]]
                data_rows = [[c.strip() for c in r.split('|')[1:-1]] for r in rows[2:]]
                
                # Fetch custom column widths
                col_w = table_col_widths.get(table_index, [6.27 / len(header)] * len(header))
                font_sz = Pt(8) if len(header) >= 9 else Pt(9)
                format_academic_table(doc, header, data_rows, col_w, font_size=font_sz)
                table_index += 1

        elif btype == "code":
            # Mermaid diagrams -> insert corresponding high-res rendered diagram PNG!
            lang = block["lang"]
            content = block["content"]
            if "graph" in content or "Baseline" in content:
                img_path = "results/system_architecture.png"
                caption = "Figure Architecture: High-level architectural topology comparison between the baseline perimeter-based model (Urmila & Shubham Mishra, 2025) and the proposed XAI-ZT-P2P 4-layer adaptive Zero-Trust framework."
                add_image_figure(doc, img_path, caption, width=Inches(5.9))
            elif "sequenceDiagram" in content or "Phase 1:" in content:
                img_path = "results/protocol_sequence.png"
                caption = "Figure Sequence: End-to-end sequence execution flow: UDP 5353 subnet beacon discovery, RSA-2048 handshake, continuous 14-feature sliding window telemetry, TreeSHAP scoring, and closed-loop 4-tier transport mitigation."
                add_image_figure(doc, img_path, caption, width=Inches(5.9))

        elif btype == "image":
            src = block["src"]
            # Clean up relative path: e.g. ../results/xxx.png -> results/xxx.png
            img_rel = src.replace("../", "")
            caption = block["caption"]
            add_image_figure(doc, img_rel, caption, width=Inches(5.8))

        elif btype == "display_math":
            content = block["content"]
            eq_num = eq_numbers[math_eq_index] if math_eq_index < len(eq_numbers) else ""
            add_display_equation(doc, content, eq_number=eq_num)
            math_eq_index += 1

        elif btype == "divider":
            # Section divider / thematic break
            pass

    print(f"Compilation finished. Total tables created: {table_index}, Display equations: {math_eq_index}")
    print(f"Saving compiled Word document to: {output_docx_path}...")
    doc.save(output_docx_path)
    file_size = os.path.getsize(output_docx_path)
    print(f"Document saved successfully! File size: {file_size:,} bytes ({file_size/1024:.1f} KB)")

if __name__ == "__main__":
    md_file = "paper/XAI_ZT_P2P_RESEARCH_PAPER.md"
    docx_file = "paper/XAI_ZT_P2P_RESEARCH_PAPER.docx"
    generate_paper_docx(md_file, docx_file)
